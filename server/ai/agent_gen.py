#!/usr/bin/env python3
"""AI Agent 文件生成: excel / word / ppt / pdf"""
import json
import os
import sys
import re
from datetime import datetime


def fail(msg):
    print(json.dumps({"ok": False, "error": str(msg)}, ensure_ascii=False))
    sys.exit(1)


def parse_body(body: str):
    body = (body or "").strip()
    lines = [ln.strip(" •-·\t") for ln in body.splitlines() if ln.strip()]
    if not lines:
        # 按中文句号切
        parts = [p.strip() for p in re.split(r"[。;;\n]+", body) if p.strip()]
        lines = parts or ["待补充内容"]
    # 尝试表格: a,b,c
    rows = []
    for ln in lines:
        if re.search(r"[,，\t|]", ln):
            cells = [c.strip() for c in re.split(r"[,，\t|]+", ln) if c.strip() != ""]
            if len(cells) >= 2:
                rows.append(cells)
    return lines, rows


def gen_excel(path, title, body):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    lines, rows = parse_body(body)
    wb = Workbook()
    ws = wb.active
    ws.title = "数据"
    header_fill = PatternFill("solid", fgColor="F2F2F2")
    header_font = Font(name="微软雅黑", bold=True, color="1F1F1F", size=12)
    body_font = Font(name="微软雅黑", size=11)
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    ws["A1"] = title
    ws["A1"].font = Font(name="微软雅黑", bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(3, len(rows[0]) if rows else 3))
    ws["A2"] = f"AI 生成 · {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws["A2"].font = Font(name="微软雅黑", size=9, color="666666")

    start = 4
    if rows:
        headers = [f"列{i+1}" for i in range(max(len(r) for r in rows))]
        if rows and all(not re.search(r"\d", rows[0][0]) for _ in [0]):
            # 首行像表头则用作表头
            if len(rows[0]) >= 2:
                headers = rows[0]
                data_rows = rows[1:] or rows
            else:
                data_rows = rows
        else:
            data_rows = rows
        for c, h in enumerate(headers, 1):
            cell = ws.cell(start, c, h)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = thin
            cell.alignment = Alignment(horizontal="center")
        for r_i, row in enumerate(data_rows, 1):
            for c_i in range(len(headers)):
                val = row[c_i] if c_i < len(row) else ""
                cell = ws.cell(start + r_i, c_i + 1, val)
                cell.font = body_font
                cell.border = thin
        for c in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + c) if c <= 26 else "A"].width = 18
    else:
        ws.cell(start, 1, "序号").fill = header_fill
        ws.cell(start, 1).font = header_font
        ws.cell(start, 2, "内容").fill = header_fill
        ws.cell(start, 2).font = header_font
        ws.cell(start, 1).border = thin
        ws.cell(start, 2).border = thin
        for i, ln in enumerate(lines, 1):
            ws.cell(start + i, 1, i).font = body_font
            ws.cell(start + i, 1).border = thin
            ws.cell(start + i, 2, ln).font = body_font
            ws.cell(start + i, 2).border = thin
        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 48
    ws.freeze_panes = "A5"
    wb.save(path)


def gen_word(path, title, body):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    lines, _ = parse_body(body)
    doc = Document()
    styles = doc.styles
    try:
        styles["Normal"].font.name = "微软雅黑"
        styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        styles["Normal"].font.size = Pt(11)
    except Exception:
        pass

    t = doc.add_heading(title, level=0)
    for run in t.runs:
        try:
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        except Exception:
            pass

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = meta.add_run(f"AI 生成文档 · {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.add_heading("内容要点", level=1)
    for ln in lines:
        p = doc.add_paragraph(ln, style="List Bullet")

    doc.add_heading("补充说明", level=1)
    doc.add_paragraph("本文档由群聊 AI 助手根据指令自动生成，可在本地用 Word / WPS 打开编辑。")
    doc.save(path)


def _set_run_cjk(run, font_name="微软雅黑"):
    """兼容旧调用"""
    from lxml import etree
    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    latin = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
    if latin is None:
        latin = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
    latin.set("typeface", font_name)
    ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    if ea is None:
        ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
    ea.set("typeface", font_name)


def gen_ppt(path, title, body):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
    from lxml import etree

    def set_run_cjk(run, font_name="微软雅黑"):
        run.font.name = font_name
        rPr = run._r.get_or_add_rPr()
        latin = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
        if latin is None:
            latin = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
        latin.set("typeface", font_name)
        ea = rPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        if ea is None:
            ea = etree.SubElement(rPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}ea")
        ea.set("typeface", font_name)

    lines, _ = parse_body(body)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def add_title_slide():
        layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.15))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0x07, 0xC1, 0x60)
        shape.line.fill.background()
        box = slide.shapes.add_textbox(Inches(0.8), Inches(2.6), Inches(11.5), Inches(1.4))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(40)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        p.alignment = PP_ALIGN.LEFT
        for run in p.runs:
            set_run_cjk(run)
        sub = slide.shapes.add_textbox(Inches(0.8), Inches(4.2), Inches(11), Inches(0.6))
        sp = sub.text_frame.paragraphs[0]
        sp.text = f"AI 助手生成 · {datetime.now().strftime('%Y-%m-%d')}"
        sp.font.size = Pt(16)
        sp.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        for run in sp.runs:
            set_run_cjk(run)

    def add_bullet_slide(idx, heading, items):
        layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        box = slide.shapes.add_textbox(Inches(0.7), Inches(0.45), Inches(12), Inches(0.8))
        p = box.text_frame.paragraphs[0]
        p.text = f"{idx}. {heading}"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        for run in p.runs:
            set_run_cjk(run)
        body_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(11.2), Inches(5.2))
        tf = body_box.text_frame
        tf.word_wrap = True
        for i, item in enumerate(items):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.text = f"• {item}"
            para.font.size = Pt(20)
            para.font.color.rgb = RGBColor(0x2E, 0x2A, 0x26)
            para.space_after = Pt(12)
            for run in para.runs:
                set_run_cjk(run)

    add_title_slide()
    chunks = [lines[i:i + 4] for i in range(0, len(lines), 4)] or [["内容生成中"]]
    for i, ch in enumerate(chunks, 1):
        add_bullet_slide(i, title if len(chunks) == 1 else f"{title} · 要点 {i}", ch)
    layout = prs.slide_layouts[6]
    end = prs.slides.add_slide(layout)
    box = end.shapes.add_textbox(Inches(0.8), Inches(2.8), Inches(11.5), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.text = "谢谢"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0x07, 0xC1, 0x60)
    for run in p.runs:
        set_run_cjk(run)
    prs.save(path)


def resolve_cjk_font():
    candidates = [
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\msyh.ttf",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc",
        "/System/Library/Fonts/PingFang.ttc",
    ]
    for c in candidates:
        if os.path.exists(c):
            return c
    return None


def gen_pdf(path, title, body):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.pdfgen import canvas
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.colors import HexColor

    lines, _ = parse_body(body)
    font_path = resolve_cjk_font()
    font_name = "Helvetica"
    if font_path:
        try:
            pdfmetrics.registerFont(TTFont("CJK", font_path, subfontIndex=0))
            font_name = "CJK"
        except Exception:
            try:
                pdfmetrics.registerFont(TTFont("CJK", font_path))
                font_name = "CJK"
            except Exception:
                font_name = "Helvetica"

    c = canvas.Canvas(path, pagesize=A4)
    w, h = A4
    y = h - 2.2 * cm
    c.setFont(font_name, 18)
    c.setFillColor(HexColor("#1F1F1F"))
    c.drawString(2.2 * cm, y, title[:40])
    y -= 0.7 * cm
    c.setFont(font_name, 9)
    c.setFillColor(HexColor("#666666"))
    c.drawString(2.2 * cm, y, f"AI 生成 · {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    y -= 0.9 * cm
    c.setStrokeColor(HexColor("#07C160"))
    c.setLineWidth(1.2)
    c.line(2.2 * cm, y, w - 2.2 * cm, y)
    y -= 0.8 * cm
    c.setFont(font_name, 12)
    c.setFillColor(HexColor("#2E2A26"))
    for ln in lines:
        if y < 2.2 * cm:
            c.showPage()
            y = h - 2.2 * cm
            c.setFont(font_name, 12)
            c.setFillColor(HexColor("#2E2A26"))
        # wrap long lines
        max_chars = 36
        while ln:
            piece, ln = ln[:max_chars], ln[max_chars:]
            c.drawString(2.2 * cm, y, piece)
            y -= 0.65 * cm
    c.save()


def main():
    try:
        data = json.load(sys.stdin)
    except Exception as e:
        fail(f"输入解析失败: {e}")
    kind = data.get("kind") or "excel"
    title = data.get("title") or "AI 文档"
    body = data.get("body") or ""
    out = data.get("out")
    if not out:
        fail("缺少输出路径")
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    try:
        if kind == "excel":
            gen_excel(out, title, body)
        elif kind == "word":
            gen_word(out, title, body)
        elif kind == "ppt":
            gen_ppt(out, title, body)
        elif kind == "pdf":
            gen_pdf(out, title, body)
        else:
            gen_excel(out, title, body)
        print(json.dumps({"ok": True, "path": out, "kind": kind}, ensure_ascii=False))
    except Exception as e:
        fail(str(e))


if __name__ == "__main__":
    main()
